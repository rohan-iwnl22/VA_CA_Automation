"""Tests for the /api/final-report endpoint — normal and TextJoin inputs."""
import io

import openpyxl
from fastapi.testclient import TestClient
from jose import jwt

from va_ca_automation.api.deps import ALGORITHM, JWT_SECRET
from va_ca_automation.api.main import create_app
from va_ca_automation.api.routes.final_report import _split_hosts

app = create_app()
client = TestClient(app)
TOKEN = jwt.encode({"sub": "admin", "role": "admin"}, JWT_SECRET, algorithm=ALGORITHM)
HEADERS = {"Authorization": f"Bearer {TOKEN}"}

HEADER_NAMES = [
    "Sr. no", "Vulnerability Title", "Description", "Risk",
    "Host", "Port", "Recommendation ", "Reference", "CVE", "Retest Status",
]

# Summary risk-table cell positions (written by _write_retest_and_summary)
RISK_ROWS = {"Critical": 16, "High": 17, "Medium": 18, "Low": 19}
GRAND_TOTAL_ROW = 20


def make_workbook(rows):
    """Build an in-memory .xlsx with headers at row 13 and data from row 14."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VA Report"
    for i, h in enumerate(HEADER_NAMES, 1):
        ws.cell(row=13, column=i, value=h)
    for r, row in enumerate(rows, 14):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()
    return buf


def post_final(first_rows, retest_rows):
    response = client.post(
        "/api/final-report",
        headers=HEADERS,
        files={
            "first_audit": (
                "first.xlsx", make_workbook(first_rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "retest_file": (
                "retest.xlsx", make_workbook(retest_rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )
    assert response.status_code == 200, response.text
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    return wb["VA Report"], wb["Summary"]


def va_rows(ws):
    """Return [(host, title, status)] for all data rows of the VA Report sheet."""
    out = []
    for row in range(14, ws.max_row + 1):
        title = ws.cell(row=row, column=2).value
        if title is None:
            continue
        out.append((
            ws.cell(row=row, column=5).value,
            title,
            ws.cell(row=row, column=10).value,
        ))
    return out


def risk_counts(summary_ws):
    return {
        label: summary_ws.cell(row=r, column=6).value
        for label, r in RISK_ROWS.items()
    }


def test_split_hosts():
    assert _split_hosts(None) == []
    assert _split_hosts("") == []
    assert _split_hosts("10.0.0.1") == ["10.0.0.1"]
    assert _split_hosts("10.0.0.1, 10.0.0.2") == ["10.0.0.1", "10.0.0.2"]
    assert _split_hosts("10.0.0.1\n10.0.0.2") == ["10.0.0.1", "10.0.0.2"]
    assert _split_hosts("10.0.0.1,") == ["10.0.0.1"]
    assert _split_hosts(" 10.0.0.1 , 10.0.0.1 , 10.0.0.2 ") == ["10.0.0.1", "10.0.0.2"]


def test_normal_inputs():
    """Baseline: one host per row — statuses, risk counts, grand total."""
    first = [
        [1, "Weak TLS", "Desc", "High", "10.0.0.1", "443", "Fix", "r1", "CVE-1", ""],
        [2, "Outdated Apache", "Desc", "Critical", "10.0.0.2", "80", "Fix", "r2", "CVE-2", ""],
        [3, "SMB Vulnerability", "Desc", "High", "10.0.0.3", "445", "Fix", "r3", "CVE-3", ""],
    ]
    retest = [
        [1, "Weak TLS", "Desc", "High", "10.0.0.1", "443", "Fix", "r1", "CVE-1", ""],
        [3, "SMB Vulnerability", "Desc", "High", "10.0.0.3", "445", "Fix", "r3", "CVE-3", ""],
    ]
    va, summary = post_final(first, retest)

    assert va_rows(va) == [
        ("10.0.0.1", "Weak TLS", "OPEN"),
        ("10.0.0.2", "Outdated Apache", "CLOSED"),
        ("10.0.0.3", "SMB Vulnerability", "OPEN"),
    ]
    assert risk_counts(summary) == {"Critical": 0, "High": 2, "Medium": 0, "Low": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 2


def test_textjoin_first_audit():
    """Joined host cells: row is OPEN if any IP is in the retest; cell stays joined."""
    first = [
        [1, "Weak TLS", "Desc", "High", "10.0.0.5, 10.0.0.6", "443", "Fix", "r1", "CVE-1", ""],
        [2, "Old SSH", "Desc", "Medium", "10.0.0.7, 10.0.0.8", "22", "Fix", "r2", "CVE-2", ""],
        [3, "SMB Vulnerability", "Desc", "High", "10.0.0.3", "445", "Fix", "r3", "CVE-3", ""],
        [4, "Newline Hosts", "Desc", "Critical", "10.0.0.9\n10.0.0.10", "80", "Fix", "r4", "CVE-4", ""],
    ]
    retest = [
        [1, "Weak TLS", "Desc", "High", "10.0.0.5", "443", "Fix", "r1", "CVE-1", ""],
        [3, "SMB Vulnerability", "Desc", "High", "10.0.0.3", "445", "Fix", "r3", "CVE-3", ""],
    ]
    va, summary = post_final(first, retest)

    rows = va_rows(va)
    assert rows[0] == ("10.0.0.5, 10.0.0.6", "Weak TLS", "OPEN")
    assert rows[1] == ("10.0.0.7, 10.0.0.8", "Old SSH", "CLOSED")
    assert rows[2] == ("10.0.0.3", "SMB Vulnerability", "OPEN")
    assert rows[3][2] == "CLOSED"
    # open IPs: 10.0.0.5 (High) + 10.0.0.3 (High)
    assert risk_counts(summary) == {"Critical": 0, "High": 2, "Medium": 0, "Low": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 2


def test_textjoin_retest_matches_single_host_rows():
    """Exploded retest lookup: joined retest cell matches a single-host audit row."""
    first = [
        [1, "SMB Vulnerability", "Desc", "High", "10.0.0.3", "445", "Fix", "r3", "CVE-3", ""],
        [2, "Outdated Apache", "Desc", "Critical", "10.0.0.2", "80", "Fix", "r2", "CVE-2", ""],
    ]
    retest = [
        [1, "SMB Vulnerability", "Desc", "High", "10.0.0.3, 10.0.0.9", "445", "Fix", "r3", "CVE-3", ""],
    ]
    va, summary = post_final(first, retest)

    assert va_rows(va) == [
        ("10.0.0.3", "SMB Vulnerability", "OPEN"),
        ("10.0.0.2", "Outdated Apache", "CLOSED"),
    ]
    assert risk_counts(summary) == {"Critical": 0, "High": 1, "Medium": 0, "Low": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 1


def test_partial_reopen_counts_only_matching_ips():
    """Only the IPs still present in the retest are counted as open."""
    first = [
        [1, "Weak TLS", "Desc", "High", "10.0.0.5, 10.0.0.6, 10.0.0.7", "443", "Fix", "r1", "CVE-1", ""],
    ]
    retest = [
        [1, "Weak TLS", "Desc", "High", "10.0.0.6", "443", "Fix", "r1", "CVE-1", ""],
    ]
    va, summary = post_final(first, retest)

    assert va_rows(va) == [("10.0.0.5, 10.0.0.6, 10.0.0.7", "Weak TLS", "OPEN")]
    # 2 of 3 IPs fixed → 1 open IP
    assert risk_counts(summary) == {"Critical": 0, "High": 1, "Medium": 0, "Low": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 1


if __name__ == "__main__":
    test_split_hosts()
    test_normal_inputs()
    test_textjoin_first_audit()
    test_textjoin_retest_matches_single_host_rows()
    test_partial_reopen_counts_only_matching_ips()
    print("All tests passed.")
