"""Tests for the /api/final-report-ca endpoint — normal and TextJoin inputs."""
import io
import shutil

import openpyxl
import pytest
from fastapi.testclient import TestClient
from jose import jwt

from va_ca_automation.api.deps import ALGORITHM, JWT_SECRET
from va_ca_automation.api.main import create_app
from va_ca_automation.api.routes import final_report_ca

app = create_app()
client = TestClient(app)
TOKEN = jwt.encode({"sub": "admin", "role": "admin"}, JWT_SECRET, algorithm=ALGORITHM)
HEADERS = {"Authorization": f"Bearer {TOKEN}"}

HEADER_NAMES = [
    "Sr. no", "Title", "Host", "Description", "Solution ", "Impact", "Status",
]

# First-audit header block: (label in col B, value in col C, rows 5-10)
METADATA = [
    ("Client Name:", "Acme Corp"),
    ("Security Tester:", "John Doe"),
    ("Reviewed By:", "Jane Roe"),
    ("Report Date:", "2026-10-09"),
    ("Report Version:", "1.0"),
    ("Scanner:", "Nessus"),
]

# Summary status-table cell positions (written by _write_ca_retest_and_summary)
STATUS_ROWS = {"FAILED": 16, "WARNING": 17}
GRAND_TOTAL_ROW = 18


@pytest.fixture
def ca_template(tmp_path, monkeypatch):
    """Point the endpoint at a temp copy of the real Final Audit CA template."""
    source = final_report_ca._TEMPLATE_DIR / "Final Auidt CA.xlsx"
    if not source.exists():
        pytest.skip(f"Final Audit CA template missing: {source}")
    target = tmp_path / "Final Auidt CA.xlsx"
    shutil.copy2(source, target)
    monkeypatch.setattr(final_report_ca, "_FINAL_AUDIT_CA_TEMPLATE", target)
    return target


def make_workbook(rows, metadata=None):
    """Build an in-memory .xlsx with CA headers at row 13 and data from row 14."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CA_Report"
    for i, h in enumerate(HEADER_NAMES, 1):
        ws.cell(row=13, column=i, value=h)
    if metadata:
        for r, (label, value) in enumerate(metadata, 5):
            ws.cell(row=r, column=2, value=label)
            ws.cell(row=r, column=3, value=value)
    for r, row in enumerate(rows, 14):
        for c, val in enumerate(row, 1):
            ws.cell(row=r, column=c, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()
    return buf


def post_final(first_rows, retest_rows, first_metadata=None):
    response = client.post(
        "/api/final-report-ca",
        headers=HEADERS,
        files={
            "first_audit": (
                "first.xlsx", make_workbook(first_rows, first_metadata),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
            "retest_file": (
                "retest.xlsx", make_workbook(retest_rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )
    assert response.status_code == 200, response.text
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    return wb["CA_Report"], wb["Summary"]


def ca_rows(ws):
    """Return [(host, title, status)] for all data rows of the CA_Report sheet."""
    out = []
    for row in range(14, ws.max_row + 1):
        title = ws.cell(row=row, column=2).value
        if title is None:
            continue
        out.append((
            ws.cell(row=row, column=3).value,
            title,
            ws.cell(row=row, column=8).value,
        ))
    return out


def status_counts(summary_ws):
    return {
        label: summary_ws.cell(row=r, column=6).value
        for label, r in STATUS_ROWS.items()
    }


def test_retest_header_added(ca_template):
    """The Retest Status header is written into the template's empty H13."""
    ca, _ = post_final(
        [[1, "TLS Config", "10.0.0.1", "d", "s", "i", "FAILED"]],
        [],
    )
    assert ca.cell(row=13, column=8).value == "Retest Status"
    assert ca.cell(row=14, column=8).value == "CLOSED"


def test_normal_inputs(ca_template):
    """Baseline: one host per row — statuses, FAILED/WARNING counts, total."""
    first = [
        [1, "Weak TLS Config", "10.0.0.1", "Desc", "Fix", "Impact", "FAILED"],
        [2, "Outdated Apache", "10.0.0.2", "Desc", "Fix", "Impact", "WARNING"],
        [3, "SMB Signing Disabled", "10.0.0.3", "Desc", "Fix", "Impact", "FAILED"],
    ]
    retest = [
        [1, "Weak TLS Config", "10.0.0.1", "Desc", "Fix", "Impact", "FAILED"],
        [3, "SMB Signing Disabled", "10.0.0.3", "Desc", "Fix", "Impact", "FAILED"],
    ]
    ca, summary = post_final(first, retest)

    assert ca_rows(ca) == [
        ("10.0.0.1", "Weak TLS Config", "OPEN"),
        ("10.0.0.2", "Outdated Apache", "CLOSED"),
        ("10.0.0.3", "SMB Signing Disabled", "OPEN"),
    ]
    assert status_counts(summary) == {"FAILED": 2, "WARNING": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 2


def test_textjoin_first_audit(ca_template):
    """Joined host cells: row is OPEN if any IP is in the retest; cell stays joined."""
    first = [
        [1, "Weak TLS Config", "10.0.0.5, 10.0.0.6", "Desc", "Fix", "Impact", "FAILED"],
        [2, "Old SSH", "10.0.0.7, 10.0.0.8", "Desc", "Fix", "Impact", "WARNING"],
        [3, "SMB Signing Disabled", "10.0.0.3", "Desc", "Fix", "Impact", "FAILED"],
        [4, "Newline Hosts", "10.0.0.9\n10.0.0.10", "Desc", "Fix", "Impact", "FAILED"],
    ]
    retest = [
        [1, "Weak TLS Config", "10.0.0.5", "Desc", "Fix", "Impact", "FAILED"],
        [3, "SMB Signing Disabled", "10.0.0.3", "Desc", "Fix", "Impact", "FAILED"],
    ]
    ca, summary = post_final(first, retest)

    rows = ca_rows(ca)
    assert rows[0] == ("10.0.0.5, 10.0.0.6", "Weak TLS Config", "OPEN")
    assert rows[1] == ("10.0.0.7, 10.0.0.8", "Old SSH", "CLOSED")
    assert rows[2] == ("10.0.0.3", "SMB Signing Disabled", "OPEN")
    assert rows[3][2] == "CLOSED"
    # open IPs: 10.0.0.5 (FAILED) + 10.0.0.3 (FAILED)
    assert status_counts(summary) == {"FAILED": 2, "WARNING": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 2


def test_textjoin_retest_matches_single_host_rows(ca_template):
    """Exploded retest lookup: joined retest cell matches a single-host audit row."""
    first = [
        [1, "SMB Signing Disabled", "10.0.0.3", "Desc", "Fix", "Impact", "FAILED"],
        [2, "Outdated Apache", "10.0.0.2", "Desc", "Fix", "Impact", "WARNING"],
    ]
    retest = [
        [1, "SMB Signing Disabled", "10.0.0.3, 10.0.0.9", "Desc", "Fix", "Impact", "FAILED"],
    ]
    ca, summary = post_final(first, retest)

    assert ca_rows(ca) == [
        ("10.0.0.3", "SMB Signing Disabled", "OPEN"),
        ("10.0.0.2", "Outdated Apache", "CLOSED"),
    ]
    assert status_counts(summary) == {"FAILED": 1, "WARNING": 0}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 1


def test_partial_reopen_counts_only_matching_ips(ca_template):
    """Only the IPs still present in the retest are counted as open."""
    first = [
        [1, "Weak TLS Config", "10.0.0.5, 10.0.0.6, 10.0.0.7",
         "Desc", "Fix", "Impact", "WARNING"],
    ]
    retest = [
        [1, "Weak TLS Config", "10.0.0.6", "Desc", "Fix", "Impact", "WARNING"],
    ]
    ca, summary = post_final(first, retest)

    assert ca_rows(ca) == [
        ("10.0.0.5, 10.0.0.6, 10.0.0.7", "Weak TLS Config", "OPEN"),
    ]
    # 2 of 3 IPs fixed → 1 open IP
    assert status_counts(summary) == {"FAILED": 0, "WARNING": 1}
    assert summary.cell(row=GRAND_TOTAL_ROW, column=6).value == 1


def test_summary_scope_table_and_leftover_rows(ca_template):
    """Scope table lists the first-audit hosts; stale template numbering is cleared."""
    first = [
        [1, "Weak TLS Config", "10.0.0.1", "Desc", "Fix", "Impact", "FAILED"],
        [2, "Old SSH", "10.0.0.2", "Desc", "Fix", "Impact", "WARNING"],
    ]
    ca, summary = post_final(first, first)

    assert summary.cell(row=8, column=1).value == 1
    assert summary.cell(row=8, column=2).value == "10.0.0.1"
    assert summary.cell(row=9, column=1).value == 2
    assert summary.cell(row=9, column=2).value == "10.0.0.2"
    # template pre-numbers rows 8-16 — rows below our 2 hosts must be blank
    for row in range(10, 17):
        assert summary.cell(row=row, column=1).value is None


def test_metadata_header_carried_over(ca_template):
    """Client Name … Scanner from the first audit are written into the template."""
    ca, _ = post_final(
        [[1, "Weak TLS Config", "10.0.0.1", "Desc", "Fix", "Impact", "FAILED"]],
        [],
        first_metadata=METADATA,
    )
    for row, (label, value) in enumerate(METADATA, 5):
        assert ca.cell(row=row, column=2).value == label
        assert ca.cell(row=row, column=3).value == value


def test_template_filename_fallback(tmp_path, monkeypatch):
    """The correctly spelled 'Final Audit CA.xlsx' is used when the primary path is gone."""
    source = final_report_ca._TEMPLATE_DIR / "Final Auidt CA.xlsx"
    if not source.exists():
        pytest.skip(f"Final Audit CA template missing: {source}")
    shutil.copy2(source, tmp_path / "Final Audit CA.xlsx")
    monkeypatch.setattr(final_report_ca, "_TEMPLATE_DIR", tmp_path)
    monkeypatch.setattr(
        final_report_ca, "_FINAL_AUDIT_CA_TEMPLATE", tmp_path / "renamed.xlsx"
    )

    ca, _ = post_final([], [])
    assert ca.cell(row=13, column=8).value == "Retest Status"


def test_output_filename(ca_template):
    response = client.post(
        "/api/final-report-ca",
        headers=HEADERS,
        files={
            "first_audit": ("CA_Server_First.xlsx", make_workbook([]), "application/octet-stream"),
            "retest_file": ("retest.xlsx", make_workbook([]), "application/octet-stream"),
        },
    )
    assert response.status_code == 200, response.text
    assert "CA_Server_First_Final.xlsx" in response.headers["content-disposition"]


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
