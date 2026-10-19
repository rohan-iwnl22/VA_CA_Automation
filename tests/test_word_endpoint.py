"""Tests for POST /api/word — any number of Excel files, no silent drops."""
import io

import openpyxl
import pytest
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi.testclient import TestClient
from jose import jwt

from va_ca_automation.api.deps import ALGORITHM, JWT_SECRET
from va_ca_automation.api.main import create_app

app = create_app()
client = TestClient(app)
TOKEN = jwt.encode({"sub": "admin", "role": "admin"}, JWT_SECRET, algorithm=ALGORITHM)
HEADERS = {"Authorization": f"Bearer {TOKEN}"}

VA_COLS = [
    "Sr. no", "Vulnerability Title", "Description", "Risk",
    "Host", "Port", "Recommendation ", "Reference", "CVE",
]
CA_COLS = ["Sr.No.", "Title", "Host", "Description", "Solution", "Impact", "Risk"]

VA_ANCHOR = "The below table shows the detailed report of the VA scan done on assets."
CA_ANCHOR = "The below table shows the detailed report of the Compliance scan done on Assets."


def make_va(idx: int, sheet_name: str = "VA Report") -> io.BytesIO:
    """Minimal VA workbook: metadata header + one finding at row 14."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["C5"] = "Test Client"
    for i, h in enumerate(VA_COLS, 1):
        ws.cell(row=13, column=i, value=h)
    ws.cell(row=14, column=1, value=1)
    ws.cell(row=14, column=2, value=f"Vulnerability {idx}")
    ws.cell(row=14, column=3, value=f"Description {idx}")
    ws.cell(row=14, column=4, value="High")
    ws.cell(row=14, column=5, value=f"10.0.0.{idx}")
    ws.cell(row=14, column=6, value="443")
    ws.cell(row=14, column=7, value="Fix it")
    ws.cell(row=14, column=8, value="ref")
    ws.cell(row=14, column=9, value=f"CVE-2026-{idx:04d}")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()
    return buf


def make_ca(idx: int, sheet_name: str = "CA_Report") -> io.BytesIO:
    """Minimal CA workbook: headers at row 13 + one finding at row 14."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["C5"] = "Test Client"
    for i, h in enumerate(CA_COLS, 1):
        ws.cell(row=13, column=i, value=h)
    ws.cell(row=14, column=1, value=1)
    ws.cell(row=14, column=2, value=f"Compliance {idx}")
    ws.cell(row=14, column=3, value=f"10.0.1.{idx}")
    ws.cell(row=14, column=4, value=f"Description {idx}")
    ws.cell(row=14, column=5, value="Apply the patch")
    ws.cell(row=14, column=6, value=f"Impact {idx}")
    ws.cell(row=14, column=7, value="Failed")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()
    return buf


def make_va_xls(idx: int, sheet_name: str = "VA Report") -> io.BytesIO:
    """Legacy BIFF .xls workbook, like the CONFIG~1.XLS files users upload."""
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet(sheet_name)
    ws.write(4, 2, "Test Client")  # C5
    for i, h in enumerate(VA_COLS):
        ws.write(12, i, h)
    ws.write(13, 0, 1)
    ws.write(13, 1, f"Vulnerability {idx}")
    ws.write(13, 2, f"Description {idx}")
    ws.write(13, 3, "High")
    ws.write(13, 4, f"10.0.0.{idx}")
    ws.write(13, 5, "443")
    ws.write(13, 6, "Fix it")
    ws.write(13, 7, "ref")
    ws.write(13, 8, f"CVE-2026-{idx:04d}")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def make_ca_rows(
    idx: int,
    findings: int,
    header_row: int = 13,
    blank_sr: bool = False,
    sheet_name: str = "CA_Report",
) -> io.BytesIO:
    """CA workbook with a configurable header position and blank Sr. no."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws["C5"] = "Test Client"
    for offset in range(3):
        ws.cell(row=1 + offset, column=1, value=f"Cover note {offset}")
    for i, h in enumerate(CA_COLS, 1):
        ws.cell(row=header_row, column=i, value=h)
    for n in range(findings):
        row = header_row + 1 + n
        if not blank_sr:
            ws.cell(row=row, column=1, value=n + 1)
        ws.cell(row=row, column=2, value=f"Compliance {idx}-{n}")
        ws.cell(row=row, column=3, value=f"10.0.1.{idx}")
        ws.cell(row=row, column=4, value=f"Description {n}")
        ws.cell(row=row, column=5, value="Apply the patch")
        ws.cell(row=row, column=6, value=f"Impact {n}")
        ws.cell(row=row, column=7, value="Failed")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()
    return buf


def make_decoy() -> io.BytesIO:
    """A workbook that is neither VA nor CA."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(row=1, column=1, value="nothing useful here")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    wb.close()
    return buf


def post_word(named_buffers) -> tuple[dict, bytes]:
    files = [
        ("files", (name, buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        for name, buf in named_buffers
    ]
    response = client.post("/api/word", headers=HEADERS, files=files, data={"client_name": "Test Client"})
    assert response.status_code == 200, response.text
    payload = response.json()
    session_id = payload["session_id"]
    download = client.get(f"/api/download/{session_id}/word", headers=HEADERS)
    assert download.status_code == 200, download.text
    return payload, download.content


def detail_tables(content: bytes) -> list:
    """VA and CA detail tables (in document order) with their kind."""
    out = []
    for table in Document(io.BytesIO(content)).tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header[:3] == ["Sr. no", "Vulnerability Title", "Description"]:
            out.append(("va", table))
        elif header[:2] == ["Sr.No.", "Title"]:
            out.append(("ca", table))
    return out


def body_sequence(content: bytes) -> list:
    """Document-order list of ('p', text) and ('va_table'|'ca_table',)."""
    doc = Document(io.BytesIO(content))
    seq = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            seq.append(("p", Paragraph(child, doc).text.strip()))
        elif tag == "tbl":
            header = [c.text.strip() for c in Table(child, doc).rows[0].cells]
            if header[:3] == ["Sr. no", "Vulnerability Title", "Description"]:
                seq.append(("va_table",))
            elif header[:2] == ["Sr.No.", "Title"]:
                seq.append(("ca_table",))
    return seq


def test_word_filename_uses_excel_client_name():
    """Filename follows VAPT_First_Audit_Report_<Client>_SCPL_<Year>_V<Version>.docx
    and the client name from Excel (C5) wins over the form value."""
    files = [("files", ("va.xlsx", make_va(1), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))]
    response = client.post(
        "/api/word",
        headers=HEADERS,
        files=files,
        data={"client_name": "Form Client", "report_number": "1.0", "report_date": "2026-03-05"},
    )
    assert response.status_code == 200, response.text
    session_id = response.json()["session_id"]
    download = client.get(f"/api/download/{session_id}/word", headers=HEADERS)
    assert download.status_code == 200, download.text
    disposition = download.headers.get("content-disposition", "")
    assert 'filename="VAPT_First_Audit_Report_Test_Client_SCPL_2026_V1_0.docx"' in disposition


def test_more_than_six_files_are_all_rendered():
    """Eight workbooks must produce eight detail tables."""
    payload, content = post_word([(f"va_scan_{i}.xlsx", make_va(i)) for i in range(1, 9)])
    assert payload["used"]["va"] == [f"va_scan_{i}.xlsx" for i in range(1, 9)]
    assert payload["used"]["ca"] == []
    assert payload["ignored"] == []
    assert len(detail_tables(content)) == 8


def test_duplicate_filenames_do_not_overwrite_each_other():
    """Two uploads sharing a filename must both survive and both be rendered."""
    _, content = post_word([("scan.xlsx", make_va(1)), ("scan.xlsx", make_va(2))])
    tables = detail_tables(content)
    assert len(tables) == 2
    titles = [t.rows[1].cells[1].text.strip() for _, t in tables]
    assert titles == ["Vulnerability 1", "Vulnerability 2"]


def test_va_and_ca_files_are_grouped_in_their_sections():
    """VA tables land in the VA section and CA tables in the CA section,
    regardless of upload order (CA files are uploaded first here)."""
    named = [(f"ca_{i}.xlsx", make_ca(i)) for i in range(1, 5)]
    named += [(f"va_{i}.xlsx", make_va(i)) for i in range(1, 5)]
    payload, content = post_word(named)

    assert payload["used"]["ca"] == [f"ca_{i}.xlsx" for i in range(1, 5)]
    assert payload["used"]["va"] == [f"va_{i}.xlsx" for i in range(1, 5)]

    seq = body_sequence(content)
    kinds = [item[0] for item in seq]
    assert kinds.count("va_table") == 4
    assert kinds.count("ca_table") == 4

    ca_anchor = next(i for i, item in enumerate(seq) if item[0] == "p" and item[1] == CA_ANCHOR)
    va_anchor = next(i for i, item in enumerate(seq) if item[0] == "p" and item[1] == VA_ANCHOR)
    va_tables = [i for i, item in enumerate(seq) if item[0] == "va_table"]
    ca_tables = [i for i, item in enumerate(seq) if item[0] == "ca_table"]

    assert va_anchor < min(va_tables)
    assert max(va_tables) < ca_anchor < min(ca_tables)


def test_renamed_sheet_is_recognised_from_headers():
    """A VA sheet renamed to 'Scan Results' is still used."""
    payload, content = post_word([("renamed.xlsx", make_va(3, sheet_name="Scan Results"))])
    assert payload["used"]["va"] == ["renamed.xlsx"]
    assert len(detail_tables(content)) == 1


def test_unrecognised_files_are_reported_not_silently_dropped():
    """A file that is neither VA nor CA shows up in 'ignored'."""
    payload, content = post_word([("va.xlsx", make_va(1)), ("notes.xlsx", make_decoy())])
    assert payload["used"]["va"] == ["va.xlsx"]
    assert len(payload["ignored"]) == 1
    assert "notes.xlsx" in payload["ignored"][0]
    assert len(detail_tables(content)) == 1


def test_legacy_xls_files_are_accepted():
    """Legacy .xls uploads (e.g. CONFIG~1.XLS) are converted and used."""
    payload, content = post_word([
        ("CONFIG~1.XLS", make_va_xls(1)),
        ("va.xlsx", make_va(2)),
    ])
    assert payload["used"]["va"] == ["CONFIG~1.XLS", "va.xlsx"]
    assert payload["ignored"] == []
    tables = detail_tables(content)
    assert len(tables) == 2
    titles = [t.rows[1].cells[1].text.strip() for _, t in tables]
    assert titles == ["Vulnerability 1", "Vulnerability 2"]


def test_three_ca_files_all_rendered_with_every_row():
    """Every CA file and every one of its finding rows must reach the report."""
    named = [
        ("ca_a.xlsx", make_ca(1, sheet_name="CA_Report")),
        ("ca_b.xlsx", make_ca(2, sheet_name="CA Report")),
        ("ca_c.xlsx", make_ca(3, sheet_name="Scan Results")),
    ]
    payload, content = post_word(named)
    assert payload["ignored"] == []
    assert payload["used"]["ca"] == ["ca_a.xlsx", "ca_b.xlsx", "ca_c.xlsx"]

    tables = [t for kind, t in detail_tables(content) if kind == "ca"]
    assert len(tables) == 3
    for idx, table in enumerate(tables, start=1):
        assert len(table.rows) == 2  # header + 1 finding
        assert table.rows[1].cells[1].text.strip() == f"Compliance {idx}"


def test_ca_rows_are_kept_even_when_layout_differs():
    """Header below row 13 and a blank Sr. no must not drop findings."""
    payload, content = post_word([
        ("offset.xlsx", make_ca_rows(1, findings=4, header_row=18)),
        ("blank_sr.xlsx", make_ca_rows(2, findings=5, blank_sr=True)),
    ])
    assert payload["ignored"] == []

    tables = [t for kind, t in detail_tables(content) if kind == "ca"]
    assert len(tables) == 2
    assert len(tables[0].rows) == 5  # header + 4 findings (header was at row 18)
    assert len(tables[1].rows) == 6  # header + 5 findings with blank Sr. no
    # Blank Sr. no is auto-numbered so the table stays readable.
    assert tables[1].rows[1].cells[0].text.strip() == "1"

