"""POST /api/final-report — Compare first audit with retest and add Retest Status."""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.drawing.line import LineProperties
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from ..deps import get_current_user

logger = logging.getLogger("va_ca_automation")

router = APIRouter(tags=["final-report"])

HEADER_ROW = 13
DATA_START_ROW = 14
SUMMARY_RETEST_HEADER_ROW = 10
SUMMARY_RETEST_DATA_ROW = 11

# Column positions in the VA Report (1-based, matching the template)
COL_TITLE = 2       # B - Vulnerability Title
COL_HOST = 5        # E - Host
COL_RETEST = 10     # J - Retest Status

# Known column name patterns for auto-detection (case-insensitive, stripped).
_TITLE_PATTERNS = re.compile(r"vulnerab.*title|title", re.IGNORECASE)
_HOST_PATTERNS = re.compile(r"\bhost\b", re.IGNORECASE)

# ── Styling constants ──────────────────────────────────────────────
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
DATA_FONT = Font(name="Cambria", size=11)
HEADER_FONT = Font(name="Cambria", size=11, bold=True)
HEADER_FILL = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")

# Retest Status colors — bold black font
OPEN_FILL = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
CLOSED_FILL = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
RETEST_FONT = Font(name="Cambria", size=11, bold=True, color="000000")

# Alignment presets
CENTER_MIDDLE = Alignment(horizontal="center", vertical="center")
LEFT_MIDDLE_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
TOP_LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
TOP_LEFT = Alignment(vertical="top", wrap_text=True)

# Per-column alignment map (1-based column index → alignment)
_COL_ALIGN = {
    1: CENTER_MIDDLE,       # Sr. no
    2: LEFT_MIDDLE_WRAP,    # Vulnerability Title — left, middle, wrap
    3: TOP_LEFT_WRAP,       # Description — top, wrap
    4: CENTER_MIDDLE,       # Risk
    5: CENTER_MIDDLE,       # Host
    6: CENTER_MIDDLE,       # Port
    7: TOP_LEFT_WRAP,       # Recommendation — top, wrap
    8: TOP_LEFT_WRAP,       # Reference — top, wrap
    9: CENTER_MIDDLE,       # CVE
    10: CENTER_MIDDLE,      # Retest Status
}

# Pie chart colors for risk levels
RISK_COLORS = {
    "Critical": "C00000",
    "High": "FF0000",
    "Medium": "FFC000",
    "Low": "FFFF00",
}

# Template path (project root / templates / final_audit_template.xlsx)
_TEMPLATE_DIR = Path(__file__).resolve().parent.parent.parent.parent.parent / "templates"
_FINAL_AUDIT_TEMPLATE = _TEMPLATE_DIR / "final_audit_template.xlsx"


def _normalize(text) -> str:
    """Lowercase, collapse whitespace, strip."""
    if text is None:
        return ""
    s = str(text).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s


def _detect_columns(ws, header_row: int = HEADER_ROW) -> dict:
    """Scan the header row to find column indices (1-based).

    Returns {"title_col", "host_col"}.
    Raises ValueError if required columns cannot be found.
    """
    title_col = None
    host_col = None
    max_col = ws.max_column or 1

    for col in range(1, max_col + 1):
        raw = ws.cell(row=header_row, column=col).value
        if raw is None:
            continue
        header = str(raw).strip()

        if title_col is None and _TITLE_PATTERNS.search(header):
            title_col = col
        if host_col is None and _HOST_PATTERNS.search(header):
            host_col = col

    if title_col is None:
        raise ValueError(
            "Could not find a Vulnerability Title / Title column in the "
            f"report header (row {header_row}).  "
            "Ensure the report has a column named 'Title', 'Vulnerability Title', "
            "or similar."
        )
    if host_col is None:
        raise ValueError(
            "Could not find a Host column in the report header "
            f"(row {header_row}).  "
            "Ensure the report has a column named 'Host'."
        )

    return {"title_col": title_col, "host_col": host_col}


def _build_retest_lookup(ws, cols: dict) -> set[tuple[str, str]]:
    """Build a normalised set of (title, host) tuples from the retest worksheet."""
    lookup: set[tuple[str, str]] = set()
    for row in range(DATA_START_ROW, ws.max_row + 1):
        title = ws.cell(row=row, column=cols["title_col"]).value
        host = ws.cell(row=row, column=cols["host_col"]).value
        if title is not None and host is not None:
            lookup.add((_normalize(title), _normalize(host)))
    return lookup


def _apply_data_cell_style(cell, col: int) -> None:
    """Apply font, border, and per-column alignment to a data cell."""
    cell.font = DATA_FONT
    cell.border = THIN_BORDER
    cell.alignment = _COL_ALIGN.get(col, CENTER_MIDDLE)


def _write_retest_and_summary(
    template_ws, summary_ws, retest_lookup: set, source_ws
) -> dict:
    """Copy data from source, write Retest Status in column J, and populate Summary.

    Returns {"open_count", "closed_count", "total", "open_by_risk": dict}.
    """
    # ── Copy data rows from source to template and determine status ──
    open_count = 0
    closed_count = 0
    total = 0
    open_by_risk: dict[str, int] = {"Critical": 0, "High": 0, "Medium": 0, "Low": 0}

    for src_row in range(DATA_START_ROW, source_ws.max_row + 1):
        # Check if row has any data
        has_data = False
        for col in range(1, 10):
            if source_ws.cell(row=src_row, column=col).value is not None:
                has_data = True
                break
        if not has_data:
            continue

        total += 1
        dest_row = DATA_START_ROW + (total - 1)

        # Copy columns A-I from source and style each cell
        for col in range(1, 10):
            cell = template_ws.cell(row=dest_row, column=col)
            cell.value = source_ws.cell(row=src_row, column=col).value
            _apply_data_cell_style(cell, col)

        # Determine Retest Status
        title = source_ws.cell(row=src_row, column=COL_TITLE).value
        host = source_ws.cell(row=src_row, column=COL_HOST).value

        if title is not None and host is not None:
            key = (_normalize(title), _normalize(host))
            status = "OPEN" if key in retest_lookup else "CLOSED"
        else:
            status = "CLOSED"

        # Write Retest Status in column J
        status_cell = template_ws.cell(row=dest_row, column=COL_RETEST)
        status_cell.value = status
        status_cell.font = RETEST_FONT
        status_cell.border = THIN_BORDER
        status_cell.alignment = CENTER_MIDDLE

        if status == "OPEN":
            status_cell.fill = OPEN_FILL
            open_count += 1
            # Count by risk level for OPEN only
            risk = source_ws.cell(row=src_row, column=4).value
            if risk:
                risk_str = str(risk).strip()
                if risk_str in open_by_risk:
                    open_by_risk[risk_str] += 1
        else:
            status_cell.fill = CLOSED_FILL
            closed_count += 1

        # Set row height
        template_ws.row_dimensions[dest_row].height = 110

    # ── Populate Summary: Scope table (IPs from VA data) ──
    hosts_seen: list[str] = []
    for src_row in range(DATA_START_ROW, source_ws.max_row + 1):
        host_val = source_ws.cell(row=src_row, column=COL_HOST).value
        if host_val is not None:
            host_str = str(host_val).strip()
            if host_str and host_str not in hosts_seen:
                hosts_seen.append(host_str)

    for i, ip in enumerate(hosts_seen):
        excel_row = 8 + i
        summary_ws.cell(row=excel_row, column=1, value=i + 1)
        summary_ws.cell(row=excel_row, column=2, value=ip)
        summary_ws.cell(row=excel_row, column=3, value="")
        for col in range(1, 4):
            cell = summary_ws.cell(row=excel_row, column=col)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_MIDDLE

    # ── Populate Summary: Risk summary (OPEN only) + Pie chart ──
    risk_header_row = 15
    summary_ws.cell(row=risk_header_row, column=5, value="Row Labels")
    summary_ws.cell(row=risk_header_row, column=6, value="Count of Host")
    for col in [5, 6]:
        cell = summary_ws.cell(row=risk_header_row, column=col)
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_MIDDLE

    # Risk data rows (only OPEN counts)
    risk_order = ["Critical", "High", "Medium", "Low"]
    data_start = risk_header_row + 1
    for i, risk in enumerate(risk_order):
        count = open_by_risk.get(risk, 0)
        summary_ws.cell(row=data_start + i, column=5, value=risk)
        summary_ws.cell(row=data_start + i, column=6, value=count)
        for col in [5, 6]:
            cell = summary_ws.cell(row=data_start + i, column=col)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_MIDDLE

    # Grand Total row
    grand_row = data_start + len(risk_order)
    summary_ws.cell(row=grand_row, column=5, value="Grand Total")
    summary_ws.cell(row=grand_row, column=6, value=open_count)
    for col in [5, 6]:
        cell = summary_ws.cell(row=grand_row, column=col)
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = CENTER_MIDDLE

    # ── Build Pie Chart ──
    chart = PieChart()
    chart.title = "Vulnerability Risk Distribution (OPEN)"
    chart.style = 10
    chart.width = 9.4
    chart.height = 9.87

    # Data source: column F (Count of Host), rows 16-19
    data_ref = Reference(summary_ws, min_col=6, min_row=data_start, max_row=data_start + len(risk_order) - 1)
    cats_ref = Reference(summary_ws, min_col=5, min_row=data_start, max_row=data_start + len(risk_order) - 1)

    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(cats_ref)

    # Apply custom colors to each slice
    series = chart.series[0]
    for idx, risk in enumerate(risk_order):
        color = RISK_COLORS.get(risk)
        if color:
            pt = DataPoint(idx=idx)
            pt.graphicalProperties.solidFill = color
            if pt.graphicalProperties.line is None:
                pt.graphicalProperties.line = LineProperties()
            pt.graphicalProperties.line.solidFill = "000000"
            pt.graphicalProperties.line.width = 10000
            series.data_points.append(pt)

    # Data labels: only value, bold font
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = False
    chart.dataLabels.showVal = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.txPr = RichText(
        p=[Paragraph(
            pPr=ParagraphProperties(
                defRPr=CharacterProperties(b=True, sz=1200, solidFill="000000")
            ),
            endParaRPr=CharacterProperties(b=True, sz=1200, solidFill="000000"),
        )]
    )

    summary_ws.add_chart(chart, "H8")

    return {
        "open_count": open_count,
        "closed_count": closed_count,
        "total": total,
        "open_by_risk": open_by_risk,
    }


@router.post("/final-report")
async def generate_final_report(
    first_audit: UploadFile = File(...),
    retest_file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
):
    """Compare first audit with retest file and generate Final Audit Report.

    Uses the final_audit_template.xlsx as base. Copies all findings from the
    first audit, adds Retest Status (OPEN/CLOSED) in column J, and populates
    the Summary sheet with scope table, risk summary (OPEN only), and pie chart.

    * **OPEN** — vulnerability + host found in retest file.
    * **CLOSED** — vulnerability in first audit but absent from retest.
    """
    # ── Validate template exists ──────────────────────────────────
    if not _FINAL_AUDIT_TEMPLATE.exists():
        raise HTTPException(
            status_code=500,
            detail=f"Final audit template not found at {_FINAL_AUDIT_TEMPLATE}",
        )

    # ── Read uploads ────────────────────────────────────────────────
    first_content = await first_audit.read()
    retest_content = await retest_file.read()

    try:
        first_wb = load_workbook(io.BytesIO(first_content))
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                f"'{first_audit.filename}' is not a valid Excel file.  "
                "Please upload a .xlsx workbook."
            ),
        )

    try:
        retest_wb = load_workbook(io.BytesIO(retest_content))
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                f"'{retest_file.filename}' is not a valid Excel file.  "
                "Please upload a .xlsx workbook."
            ),
        )

    # ── Detect columns from the first audit ─────────────────────────
    first_ws = first_wb.active
    try:
        first_cols = _detect_columns(first_ws)
    except ValueError as exc:
        first_wb.close()
        retest_wb.close()
        raise HTTPException(status_code=400, detail=str(exc))

    # ── Detect columns from the retest file ─────────────────────────
    retest_ws = retest_wb.active
    try:
        retest_cols = _detect_columns(retest_ws)
    except ValueError as exc:
        first_wb.close()
        retest_wb.close()
        raise HTTPException(status_code=400, detail=str(exc))

    # ── Build lookup from retest file ───────────────────────────────
    retest_lookup = _build_retest_lookup(retest_ws, retest_cols)
    logger.info(
        "Final report: retest lookup contains %d unique (title, host) pairs",
        len(retest_lookup),
    )

    # ── Load template ───────────────────────────────────────────────
    template_wb = load_workbook(_FINAL_AUDIT_TEMPLATE)
    template_ws = template_wb["VA Report"]
    summary_ws = template_wb["Summary"]

    # ── Write data, retest status, and summary ──────────────────────
    stats = _write_retest_and_summary(template_ws, summary_ws, retest_lookup, first_ws)
    logger.info(
        "Final report: %d OPEN, %d CLOSED out of %d total | OPEN by risk: %s",
        stats["open_count"],
        stats["closed_count"],
        stats["total"],
        stats["open_by_risk"],
    )

    # ── Save to buffer ──────────────────────────────────────────────
    buf = io.BytesIO()
    template_wb.save(buf)
    buf.seek(0)
    template_wb.close()
    first_wb.close()
    retest_wb.close()

    # ── Output filename ─────────────────────────────────────────────
    base_name = "Final_Audit_Report"
    if first_audit.filename:
        stem = first_audit.filename.rsplit(".", 1)[0]
        base_name = f"{stem}_Final"

    output_name = f"{base_name}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{output_name}"'},
    )
