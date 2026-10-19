"""POST /api/final-report-ca — Compare CA first audit with retest and add Retest Status.

Same logic as the VA final report (`/api/final-report`), adapted to the CA
layout: findings are copied from the CA first audit into the
``Final Audit CA.xlsx`` template (sheet ``CA_Report``), the first audit's
header block is carried over, Retest Status is written in column H, and the
Summary sheet gets the scope table, an OPEN-only status summary
(FAILED / WARNING) and a pie chart.

Normal (one host per row) and TextJoin (comma-joined hosts) inputs are
supported in any combination — host cells are always checked per IP.
"""

from __future__ import annotations

import io
import logging
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.styles import Alignment
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from ..deps import get_current_user
from .final_report import (
    CLOSED_FILL,
    DATA_FONT,
    DATA_START_ROW,
    HEADER_FILL,
    HEADER_FONT,
    OPEN_FILL,
    RETEST_FONT,
    THIN_BORDER,
    _build_retest_lookup,
    _detect_columns,
    _normalize,
    _split_hosts,
)

logger = logging.getLogger("va_ca_automation")

router = APIRouter(tags=["final-report"])

# CA layout in the Final Audit CA template (1-based columns)
COL_CA_TITLE = 2        # B - Title
COL_CA_HOST = 3         # C - Host
COL_CA_STATUS_RISK = 7  # G - Status (FAILED / WARNING)
COL_CA_RETEST = 8       # H - Retest Status
CA_DATA_LAST_COL = 7    # A-G are copied verbatim from the first audit

# CA status categories used in the Summary risk table
CA_STATUS_ORDER = ["FAILED", "WARNING"]

# Pie chart slice colors — match the CA report's own FAILED/WARNING fills
CA_STATUS_COLORS = {
    "FAILED": "4472C4",
    "WARNING": "ED7D31",
}

# Per-column alignment map for CA data cells (1-based)
_CA_COL_ALIGN = {
    1: Alignment(horizontal="center", vertical="center"),                 # Sr. no
    2: Alignment(horizontal="left", vertical="center", wrap_text=True),   # Title
    3: Alignment(horizontal="center", vertical="center", wrap_text=True), # Host
    4: Alignment(horizontal="left", vertical="top", wrap_text=True),      # Description
    5: Alignment(horizontal="left", vertical="top", wrap_text=True),      # Solution
    6: Alignment(horizontal="left", vertical="top", wrap_text=True),      # Impact
    7: Alignment(horizontal="center", vertical="center"),                 # Status
    8: Alignment(horizontal="center", vertical="center"),                 # Retest Status
}

CENTER_MIDDLE = Alignment(horizontal="center", vertical="center")

_TEMPLATE_DIR = Path(__file__).resolve().parent.parent.parent.parent.parent / "templates"
_FINAL_AUDIT_CA_TEMPLATE = _TEMPLATE_DIR / "Final Auidt CA.xlsx"

# Accepted template filenames — the shipped file carries the "Auidt" typo.
_TEMPLATE_NAMES = ("Final Audit CA.xlsx", "Final Auidt CA.xlsx")

# First-audit header block rows copied into the template (label in col B).
_METADATA_ROWS = range(5, 11)


def _resolve_template() -> Path:
    """Return the Final Audit CA template path.

    ``_FINAL_AUDIT_CA_TEMPLATE`` is tried first (tests point it at a temp
    copy), then either accepted filename in the templates directory.
    """
    candidates = [_FINAL_AUDIT_CA_TEMPLATE]
    for name in _TEMPLATE_NAMES:
        path = _TEMPLATE_DIR / name
        if path not in candidates:
            candidates.append(path)
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return candidates[0]


def _apply_ca_data_cell_style(cell, col: int) -> None:
    """Apply font, border, and per-column alignment to a CA data cell."""
    cell.font = DATA_FONT
    cell.border = THIN_BORDER
    cell.alignment = _CA_COL_ALIGN.get(col, CENTER_MIDDLE)


def _ensure_retest_header(template_ws) -> None:
    """Make sure the Retest Status header exists at H13.

    The Final Audit CA template ships with an empty H13; the header style is
    cloned from the neighbouring Status header (G13).
    """
    header = template_ws.cell(row=13, column=COL_CA_RETEST)
    if header.value in (None, ""):
        source = template_ws.cell(row=13, column=COL_CA_STATUS_RISK)
        header.value = "Retest Status"
        header.font = copy(source.font)
        header.fill = copy(source.fill)
        header.border = copy(source.border)
        header.alignment = copy(source.alignment)


def _report_sheet(wb):
    """Return the CA_Report sheet when present, else the workbook's active sheet."""
    if "CA_Report" in wb.sheetnames:
        return wb["CA_Report"]
    return wb.active


def _copy_metadata(source_ws, template_ws) -> None:
    """Carry the first audit's header block (Client Name … Scanner) over.

    Rows 5-10 of the CA_Report sheet hold the label in column B and the
    value in column C.  A row is only written when the label matches and
    the first audit actually has a value.
    """
    for row in _METADATA_ROWS:
        label = source_ws.cell(row=row, column=2).value
        value = source_ws.cell(row=row, column=3).value
        if value in (None, ""):
            continue
        template_label = template_ws.cell(row=row, column=2).value
        if _normalize(template_label) != _normalize(label):
            continue
        template_ws.cell(row=row, column=3, value=value)


def _normalize_status(value) -> str:
    """Normalise a CA status cell to FAILED / WARNING (or '')."""
    if value is None:
        return ""
    status = str(value).strip().upper()
    if status == "WARNINGS":
        status = "WARNING"
    return status


def _write_ca_retest_and_summary(
    template_ws, summary_ws, retest_lookup: set, source_ws, cols: dict
) -> dict:
    """Copy CA findings into the template, add Retest Status, populate Summary.

    Columns A-G are copied as-is (so TextJoin host cells stay joined).
    A row is OPEN when at least one of its hosts is found in the retest for
    the same title, otherwise CLOSED. The status is written in column H.

    Returns {"open_count", "closed_count", "total", "open_by_status",
             "open_ip_count"}.
    """
    _ensure_retest_header(template_ws)

    title_col = cols["title_col"]
    host_col = cols["host_col"]

    # Remember the scope rows pre-numbered in the template so leftovers can
    # be cleared when the new scope list is shorter.
    template_scope_rows = [
        row
        for row in range(8, summary_ws.max_row + 1)
        if summary_ws.cell(row=row, column=1).value is not None
    ]

    open_count = 0
    closed_count = 0
    total = 0
    open_ip_count = 0
    open_by_status: dict[str, int] = {status: 0 for status in CA_STATUS_ORDER}

    for src_row in range(DATA_START_ROW, source_ws.max_row + 1):
        has_data = False
        for col in range(1, CA_DATA_LAST_COL + 1):
            if source_ws.cell(row=src_row, column=col).value is not None:
                has_data = True
                break
        if not has_data:
            continue

        total += 1
        dest_row = DATA_START_ROW + (total - 1)

        # Copy columns A-G from the first audit
        for col in range(1, CA_DATA_LAST_COL + 1):
            cell = template_ws.cell(row=dest_row, column=col)
            cell.value = source_ws.cell(row=src_row, column=col).value
            _apply_ca_data_cell_style(cell, col)

        # Retest Status — per IP, so normal and TextJoin cells both work
        title = source_ws.cell(row=src_row, column=title_col).value
        host = source_ws.cell(row=src_row, column=host_col).value

        open_ips: list[str] = []
        if title is not None and host is not None:
            title_key = _normalize(title)
            open_ips = [
                ip
                for ip in _split_hosts(host)
                if (title_key, _normalize(ip)) in retest_lookup
            ]
        status = "OPEN" if open_ips else "CLOSED"

        status_cell = template_ws.cell(row=dest_row, column=COL_CA_RETEST)
        status_cell.value = status
        status_cell.font = RETEST_FONT
        status_cell.border = THIN_BORDER
        status_cell.alignment = CENTER_MIDDLE

        if status == "OPEN":
            status_cell.fill = OPEN_FILL
            open_count += 1
            open_ip_count += len(open_ips)
            ca_status = _normalize_status(
                source_ws.cell(row=src_row, column=COL_CA_STATUS_RISK).value
            )
            if ca_status in open_by_status:
                open_by_status[ca_status] += len(open_ips)
        else:
            status_cell.fill = CLOSED_FILL
            closed_count += 1

        template_ws.row_dimensions[dest_row].height = 110

    # ── Summary: scope table (host cells copied as-is) ──────────────
    hosts_seen: list[str] = []
    for src_row in range(DATA_START_ROW, source_ws.max_row + 1):
        host_val = source_ws.cell(row=src_row, column=host_col).value
        if host_val is not None:
            host_str = str(host_val).strip()
            if host_str and host_str not in hosts_seen:
                hosts_seen.append(host_str)

    for i, ip in enumerate(hosts_seen):
        excel_row = 8 + i
        summary_ws.cell(row=excel_row, column=1, value=i + 1)
        summary_ws.cell(row=excel_row, column=2, value=ip)
        summary_ws.cell(row=excel_row, column=3, value="")
        for col in range(1, 4):
            cell = summary_ws.cell(row=excel_row, column=col)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_MIDDLE

    # Clear pre-numbered template rows below the new scope list
    last_scope_row = 7 + len(hosts_seen)
    for row in template_scope_rows:
        if row > last_scope_row:
            for col in range(1, 4):
                summary_ws.cell(row=row, column=col).value = None

    # ── Summary: status summary (OPEN only) + pie chart ─────────────
    risk_header_row = 15
    summary_ws.cell(row=risk_header_row, column=5, value="Row Labels")
    summary_ws.cell(row=risk_header_row, column=6, value="Count of Host")
    for col in [5, 6]:
        cell = summary_ws.cell(row=risk_header_row, column=col)
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_MIDDLE

    data_start = risk_header_row + 1
    for i, status_label in enumerate(CA_STATUS_ORDER):
        count = open_by_status.get(status_label, 0)
        summary_ws.cell(row=data_start + i, column=5, value=status_label)
        summary_ws.cell(row=data_start + i, column=6, value=count)
        for col in [5, 6]:
            cell = summary_ws.cell(row=data_start + i, column=col)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER_MIDDLE

    grand_row = data_start + len(CA_STATUS_ORDER)
    summary_ws.cell(row=grand_row, column=5, value="Grand Total")
    summary_ws.cell(row=grand_row, column=6, value=sum(open_by_status.values()))
    for col in [5, 6]:
        cell = summary_ws.cell(row=grand_row, column=col)
        cell.font = HEADER_FONT
        cell.border = THIN_BORDER
        cell.alignment = CENTER_MIDDLE

    # ── Pie chart ───────────────────────────────────────────────────
    chart = PieChart()
    chart.title = "Compliance Risk Distribution (OPEN)"
    chart.style = 10
    chart.width = 9.4
    chart.height = 9.87

    data_ref = Reference(
        summary_ws,
        min_col=6,
        min_row=data_start,
        max_row=data_start + len(CA_STATUS_ORDER) - 1,
    )
    cats_ref = Reference(
        summary_ws,
        min_col=5,
        min_row=data_start,
        max_row=data_start + len(CA_STATUS_ORDER) - 1,
    )

    chart.add_data(data_ref, titles_from_data=False)
    chart.set_categories(cats_ref)

    series = chart.series[0]
    for idx, status_label in enumerate(CA_STATUS_ORDER):
        color = CA_STATUS_COLORS.get(status_label)
        if color:
            pt = DataPoint(idx=idx)
            pt.graphicalProperties.solidFill = color
            if pt.graphicalProperties.line is None:
                pt.graphicalProperties.line = LineProperties()
            pt.graphicalProperties.line.solidFill = "000000"
            pt.graphicalProperties.line.width = 10000
            series.data_points.append(pt)

    chart.dataLabels = DataLabelList()
    chart.dataLabels.showPercent = False
    chart.dataLabels.showVal = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.txPr = RichText(
        p=[
            Paragraph(
                pPr=ParagraphProperties(
                    defRPr=CharacterProperties(b=True, sz=1200, solidFill="000000")
                ),
                endParaRPr=CharacterProperties(b=True, sz=1200, solidFill="000000"),
            )
        ]
    )

    summary_ws.add_chart(chart, "H8")

    return {
        "open_count": open_count,
        "closed_count": closed_count,
        "total": total,
        "open_by_status": open_by_status,
        "open_ip_count": open_ip_count,
    }


@router.post("/final-report-ca")
async def generate_final_report_ca(
    first_audit: UploadFile = File(...),
    retest_file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
):
    """Compare the CA first audit with the retest file and generate the Final CA report.

    Uses ``Final Audit CA.xlsx`` as the base template. Copies all findings
    from the first audit (columns A-G, host cells as-is so TextJoin inputs
    stay joined), carries the header block (Client Name … Scanner) over,
    adds Retest Status (OPEN/CLOSED) in column H, and populates the Summary
    sheet with the scope table, an OPEN-only FAILED/WARNING summary and a
    pie chart.

    * **OPEN** — at least one IP of the row's host cell is found in the
      retest for the same title.
    * **CLOSED** — none of the row's IPs appear in the retest.
    """
    template_path = _resolve_template()
    if not template_path.exists():
        raise HTTPException(
            status_code=500,
            detail=(
                "Final Audit CA template not found — expected one of: "
                + ", ".join(str(_TEMPLATE_DIR / n) for n in _TEMPLATE_NAMES)
            ),
        )

    first_content = await first_audit.read()
    retest_content = await retest_file.read()

    try:
        first_wb = load_workbook(io.BytesIO(first_content))
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                f"'{first_audit.filename}' is not a valid Excel file.  "
                "Please upload a .xlsx workbook."
            ),
        )

    try:
        retest_wb = load_workbook(io.BytesIO(retest_content))
    except Exception:
        raise HTTPException(
            status_code=400,
            detail=(
                f"'{retest_file.filename}' is not a valid Excel file.  "
                "Please upload a .xlsx workbook."
            ),
        )

    # ── Detect columns in both uploads ──────────────────────────────
    first_ws = _report_sheet(first_wb)
    try:
        first_cols = _detect_columns(first_ws)
    except ValueError as exc:
        first_wb.close()
        retest_wb.close()
        raise HTTPException(status_code=400, detail=str(exc))

    retest_ws = _report_sheet(retest_wb)
    try:
        retest_cols = _detect_columns(retest_ws)
    except ValueError as exc:
        first_wb.close()
        retest_wb.close()
        raise HTTPException(status_code=400, detail=str(exc))

    # ── Build retest lookup (explodes TextJoin host cells) ──────────
    retest_lookup = _build_retest_lookup(retest_ws, retest_cols)
    logger.info(
        "Final CA report: retest lookup contains %d unique (title, ip) pairs",
        len(retest_lookup),
    )

    # ── Load template ───────────────────────────────────────────────
    try:
        template_wb = load_workbook(template_path)
    except PermissionError:
        raise HTTPException(
            status_code=500,
            detail=(
                f"Final Audit CA template is locked by another program: "
                f"{template_path.name}.  Close it in Excel and try again."
            ),
        )
    template_ws = template_wb["CA_Report"]
    summary_ws = template_wb["Summary"]

    _copy_metadata(first_ws, template_ws)
    stats = _write_ca_retest_and_summary(
        template_ws, summary_ws, retest_lookup, first_ws, first_cols
    )
    logger.info(
        "Final CA report: %d OPEN, %d CLOSED out of %d total (%d open IPs) "
        "| OPEN by status: %s",
        stats["open_count"],
        stats["closed_count"],
        stats["total"],
        stats["open_ip_count"],
        stats["open_by_status"],
    )

    buf = io.BytesIO()
    template_wb.save(buf)
    buf.seek(0)
    template_wb.close()
    first_wb.close()
    retest_wb.close()

    base_name = "Final_Audit_CA_Report"
    if first_audit.filename:
        stem = first_audit.filename.rsplit(".", 1)[0]
        base_name = f"{stem}_Final"

    output_name = f"{base_name}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{output_name}"'},
    )
