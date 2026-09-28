"""Final Audit Report pipeline — compare VA Normal vs Rescan."""

from __future__ import annotations

import logging
import os
from pathlib import Path

from openpyxl import load_workbook

from ..excel_writer.chart_builder import build_pie_chart
from ..excel_writer.data_writer import (
    build_rescan_lookup,
    extract_va_report_data,
    write_final_audit_rows,
    write_introduction_fields,
    write_va_report_header,
)
from ..excel_writer.summary_builder import (
    build_risk_summary,
    build_scope_table,
    write_risk_summary_table,
    write_scope_table,
)
from ..excel_writer.template_cloner import clone_template, load_working_copy
from ..logging.pipeline_logger import PipelineLogger
from ..metadata.engagement_metadata import EngagementMetadata
from ..naming.filename_builder import build_filename, ensure_unique_path

logger = logging.getLogger("va_ca_automation")


def _safe_save_workbook(wb, path: Path) -> None:
    """Save workbook, handling PermissionError by removing locked file if possible."""
    try:
        wb.save(path)
    except PermissionError:
        logger.warning("File %s is locked. Attempting to remove and retry...", path)
        try:
            os.remove(path)
            wb.save(path)
        except OSError as e:
            logger.error("Cannot remove locked file: %s", e)
            new_path = path.parent / f"{path.stem}_locked{path.suffix}"
            logger.warning("Saving to alternative path: %s", new_path)
            wb.save(new_path)


def run_final_audit_pipeline(
    va_normal_path: Path | str,
    rescan_path: Path | str,
    template_path: Path | str,
    metadata: EngagementMetadata,
    output_dir: Path | str,
    log_file: Path | None = None,
) -> Path:
    """Generate Final Audit report by comparing VA Normal vs Rescan.

    Parameters
    ----------
    va_normal_path : Path
        Path to the VA Normal report (pipeline output).
    rescan_path : Path
        Path to the Rescan report.
    template_path : Path
        Path to the pristine blank template (.xlsx).
    metadata : EngagementMetadata
        Engagement metadata for populating the report.
    output_dir : Path
        Directory where the final report will be saved.
    log_file : Path, optional
        If provided, structured log entries will be written here.

    Returns
    -------
    Path
        Path to the generated Final Audit report file.
    """
    va_normal_path = Path(va_normal_path)
    rescan_path = Path(rescan_path)
    template_path = Path(template_path)
    output_dir = Path(output_dir)

    plogger = PipelineLogger(log_file=log_file)

    # 1. LOAD both Excel files
    logger.info("Loading VA Normal report: %s", va_normal_path)
    va_normal_wb = load_workbook(va_normal_path)

    logger.info("Loading Rescan report: %s", rescan_path)
    rescan_wb = load_workbook(rescan_path)

    try:
        # 2. EXTRACT VA Report data from both files
        va_normal_ws = va_normal_wb["VA Report"]
        rescan_ws = rescan_wb["VA Report"]

        va_data = extract_va_report_data(va_normal_ws)
        plogger.log_stage_count("va_normal_rows", len(va_data))

        rescan_data = extract_va_report_data(rescan_ws)
        plogger.log_stage_count("rescan_rows", len(rescan_data))

        # 3. BUILD Rescan lookup set
        rescan_lookup = build_rescan_lookup(rescan_data)
        plogger.log_stage_count("rescan_lookup_size", len(rescan_lookup))

        # 4. CLONE template
        filename = build_filename(metadata, report_type="VA_Final_Audit")
        output_path = output_dir / filename
        output_path = ensure_unique_path(output_path)

        working_path = clone_template(template_path, output_path)
        logger.info("Working copy created: %s", working_path)

        # 5. WRITE VA Report sheet
        wb = load_working_copy(working_path)
        try:
            va_ws = wb["VA Report"]

            # Write header metadata (rows 5-10, column C)
            write_va_report_header(va_ws, metadata)

            # Write data rows with Retest Status
            last_row = write_final_audit_rows(va_ws, va_data, rescan_lookup)
            plogger.log_stage_count("final_audit_rows_written", last_row - 13)

            # 6. WRITE Summary sheet
            import pandas as pd

            va_df = pd.DataFrame(va_data)
            summary_ws = wb["Summary"]
            scope_df = build_scope_table(va_df, metadata)
            write_scope_table(summary_ws, scope_df)

            risk_summary = build_risk_summary(va_df)
            write_risk_summary_table(summary_ws, risk_summary, start_row=15)
            build_pie_chart(summary_ws, risk_summary, chart_anchor="I8", data_start_row=15)

            # 7. WRITE Introduction sheet
            intro_ws = wb["Introduction"]
            write_introduction_fields(intro_ws, metadata)

            # 8. SAVE workbook
            _safe_save_workbook(wb, working_path)
            plogger.log_output_file(str(working_path))
            logger.info("Final Audit report saved: %s", working_path)

        except Exception:
            wb.close()
            raise

    finally:
        va_normal_wb.close()
        rescan_wb.close()

    plogger.log_summary()
    plogger.flush()

    return working_path
